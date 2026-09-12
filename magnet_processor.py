# magnet_processor.py
import requests
import openpyxl
import xlsxwriter
from openpyxl import load_workbook
from PIL import Image
import io
import tempfile
import os
import time
import sys
import glob
import json
import hashlib
import subprocess
import fcntl
import threading
import re
from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

# ========== 用户可调配置 ==========
BATCH_SIZE = 50                 # 每批处理行数（处理满 50 行立即保存一个 Excel 文件）
MAX_RETRIES = 5                 # 下载图片重试次数
TIMEOUT = 20                    # 请求超时（秒）
DEFAULT_COL_WIDTH = 80          # 截图列宽（字符数）
DEFAULT_ROW_HEIGHT = 274        # 截图行高（磅）
JPEG_QUALITY = 85               # JPEG 压缩质量
MAX_RUN_TIME = 330              # 单次最大运行时间（分钟）

# ========== 并行 & 限速配置 ==========
MAX_WORKERS = int(os.environ.get('MAX_WORKERS', 4))                     # 跨行进程数
SCREENSHOT_THREADS = int(os.environ.get('SCREENSHOT_THREADS', 6))       # 行内截图并行线程数
API_RATE_LIMIT = int(os.environ.get('API_RATE_LIMIT', 5))               # 每分钟最大 API 请求数（官方限制 5）
API_RATE_WINDOW = 60                                                     # 限速时间窗口（秒）
# ======================================

# 尝试导入 PIL
try:
    from PIL import Image as PILImage
    PILLOW_AVAILABLE = True
except ImportError:
    PILLOW_AVAILABLE = False
    print("错误: 未安装 Pillow，请安装: pip install Pillow")
    sys.exit(1)


# ============================================================
# 全局限速器（跨进程安全，基于文件锁 + 时间戳滑动窗口）
# ============================================================
class GlobalRateLimiter:
    """
    跨进程共享的令牌桶限速器。
    使用文件锁（fcntl.flock）保证多进程下的原子性，
    通过 JSON 文件记录最近的请求时间戳，实现滑动窗口限速。
    """

    def __init__(self, rate=None, window=None, state_file=None):
        self.rate = rate if rate is not None else API_RATE_LIMIT
        self.window = window if window is not None else API_RATE_WINDOW
        if state_file is None:
            state_file = str(Path(tempfile.gettempdir()) / "whatslink_rate_limit.json")
        self.state_file = Path(state_file)
        self.lock_file = self.state_file.with_suffix('.lock')
        self._thread_lock = threading.Lock()

    def acquire(self):
        """阻塞直到获取一个令牌（可安全发出一次 API 请求）"""
        while True:
            sleep_time = 0.0
            with self._thread_lock:
                try:
                    with open(self.lock_file, 'w') as lf:
                        fcntl.flock(lf.fileno(), fcntl.LOCK_EX)
                        try:
                            now = time.time()
                            timestamps = []
                            if self.state_file.exists():
                                try:
                                    with open(self.state_file, 'r') as f:
                                        data = json.load(f)
                                    timestamps = [
                                        t for t in data.get('timestamps', [])
                                        if now - t < self.window
                                    ]
                                except (json.JSONDecodeError, IOError, ValueError):
                                    timestamps = []

                            if len(timestamps) < self.rate:
                                timestamps.append(now)
                                try:
                                    with open(self.state_file, 'w') as f:
                                        json.dump({'timestamps': timestamps}, f)
                                except IOError:
                                    pass
                                return
                            else:
                                sleep_time = self.window - (now - timestamps[0]) + 0.05
                        finally:
                            fcntl.flock(lf.fileno(), fcntl.LOCK_UN)
                except Exception as e:
                    print(f"  ⚠️ 限速器异常: {e}")
                    time.sleep(self.window / max(self.rate, 1))
                    return

            if sleep_time > 0:
                time.sleep(sleep_time)


# 每个进程独享的限速器实例（fork 后 PID 变化时会重建）
_rate_limiter = None
_rate_limiter_pid = None
_rate_limiter_lock = threading.Lock()


def get_rate_limiter():
    """获取当前进程的限速器单例（自动处理 fork 场景）"""
    global _rate_limiter, _rate_limiter_pid
    current_pid = os.getpid()
    with _rate_limiter_lock:
        if _rate_limiter is None or _rate_limiter_pid != current_pid:
            _rate_limiter = GlobalRateLimiter()
            _rate_limiter_pid = current_pid
        return _rate_limiter


# ============================================================
# 工具函数
# ============================================================
def format_size(size_bytes):
    """字节转可读格式"""
    if size_bytes < 1024:
        return f"{size_bytes} B"
    elif size_bytes < 1024 ** 2:
        return f"{size_bytes / 1024:.2f} KB"
    elif size_bytes < 1024 ** 3:
        return f"{size_bytes / 1024 ** 2:.2f} MB"
    elif size_bytes < 1024 ** 4:
        return f"{size_bytes / 1024 ** 3:.2f} GB"
    else:
        return f"{size_bytes / 1024 ** 4:.2f} TB"


def get_magnet_info(magnet_link):
    """调用 whatslink.info API 获取磁力信息（遵守官方 5次/分钟 限制）"""
    get_rate_limiter().acquire()  # 阻塞等待令牌，确保不超限

    url = "https://whatslink.info/api/v1/link"
    params = {"url": magnet_link}
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                      "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36"
    }
    try:
        resp = requests.get(url, params=params, headers=headers, timeout=TIMEOUT)
        resp.raise_for_status()
        data = resp.json()

        if data.get("error"):
            print(f"    API返回错误: {data['error']}")
            return None

        name = data.get("name", "").strip()
        count = data.get("count", 0)
        size_bytes = data.get("size", 0)
        size_str = format_size(size_bytes)

        screenshots = []
        for item in data.get("screenshots", []):
            if isinstance(item, dict) and item.get("screenshot"):
                screenshots.append({
                    "time": item.get("time", 0),
                    "screenshot": item.get("screenshot", "")
                })

        return {
            "name": name,
            "count": count,
            "size": size_str,
            "screenshots": screenshots
        }
    except Exception as e:
        print(f"    API请求失败: {e}")
        return None


def download_and_convert_to_jpeg(url, quality=JPEG_QUALITY, max_retries=MAX_RETRIES, timeout=TIMEOUT):
    """下载图片并转换为 JPEG，返回临时文件路径"""
    for attempt in range(max_retries):
        try:
            resp = requests.get(url, timeout=timeout)
            resp.raise_for_status()
            img = PILImage.open(io.BytesIO(resp.content))
            if img.mode == 'RGBA':
                img = img.convert('RGB')
            fd, tmp_path = tempfile.mkstemp(suffix='.jpg')
            os.close(fd)
            img.save(tmp_path, format='JPEG', quality=quality)
            return tmp_path
        except Exception as e:
            print(f"    下载/转换图片失败 (尝试 {attempt + 1}/{max_retries}): {e}")
            if attempt < max_retries - 1:
                time.sleep(2)
            else:
                return None
    return None


def download_screenshots_parallel(urls, max_threads=None):
    """
    并行下载一行内的所有截图（线程池）。
    由于 requests 是 I/O 密集型，等待网络时会释放 GIL，
    因此线程池在此场景几乎达到进程级性能，且无序列化开销。
    返回成功下载的临时文件路径列表。
    """
    if not urls:
        return []

    if max_threads is None:
        max_threads = SCREENSHOT_THREADS

    tmp_files = []
    with ThreadPoolExecutor(max_workers=max_threads) as tpe:
        future_to_url = {tpe.submit(download_and_convert_to_jpeg, u): u for u in urls}
        for future in as_completed(future_to_url):
            try:
                path = future.result()
                if path:
                    tmp_files.append(path)
            except Exception as e:
                print(f"    截图下载异常: {e}")
    return tmp_files


def get_file_hash(file_path):
    """获取文件的 MD5 哈希值"""
    hasher = hashlib.md5()
    with open(file_path, 'rb') as f:
        buf = f.read(65536)
        while len(buf) > 0:
            hasher.update(buf)
            buf = f.read(65536)
    return hasher.hexdigest()


def git_commit_files(output_dir, batch_num):
    """提交指定批次文件和进度文件到 Git（带 push 重试）"""
    try:
        subprocess.run(['git', 'config', '--global', 'user.name', 'github-actions[bot]'],
                       check=False, capture_output=True)
        subprocess.run(['git', 'config', '--global', 'user.email',
                        'github-actions[bot]@users.noreply.github.com'],
                       check=False, capture_output=True)

        progress_file = output_dir / '.progress.json'
        if progress_file.exists():
            subprocess.run(['git', 'add', '-f', str(progress_file)],
                           check=False, capture_output=True)

        batch_pattern = f'{batch_num:03d}_*.xlsx'
        for xlsx_file in glob.glob(str(output_dir / batch_pattern)):
            subprocess.run(['git', 'add', '-f', xlsx_file],
                           check=False, capture_output=True)

        result = subprocess.run(['git', 'diff', '--staged', '--quiet'], capture_output=True)

        if result.returncode != 0:
            commit_msg = f"更新进度和批次 {batch_num:03d} [skip ci]"
            subprocess.run(['git', 'commit', '-m', commit_msg],
                           check=True, capture_output=True)

            # push 带重试，防止并发场景下的远端冲突
            for _ in range(3):
                push_result = subprocess.run(['git', 'push'], capture_output=True)
                if push_result.returncode == 0:
                    print(f"  ✅ Git 提交成功: 批次 {batch_num:03d}")
                    return True
                subprocess.run(['git', 'pull', '--rebase'], capture_output=True)
                time.sleep(2)

            print(f"  ⚠️ Git push 失败（已重试）")
            return False
        else:
            print(f"  ⏭️  无变更，跳过提交")
            return False

    except Exception as e:
        print(f"  ⚠️  Git 提交失败: {e}")
        return False


# ============================================================
# 多进程 worker（必须是模块级函数以便 pickle）
# ============================================================
def _process_row_worker(args):
    """
    单个进程处理一行磁力链接：
      1. 调用 API 获取信息（通过全局限速器保证 <= 5 次/分钟）
      2. 行内并行下载并转换所有截图（线程池）
      3. 返回结果字典
    """
    magnet, row_num = args

    result = {
        'magnet': magnet,
        'name': '',
        'count': '',
        'size': '',
        'screenshots': [],
        'temp_files': [],
        'row_num': row_num,
    }

    if not magnet or not magnet.startswith("magnet:"):
        result['name'] = "Invalid Link"
        return result

    info = get_magnet_info(magnet)
    if not info:
        result['name'] = "Error"
        return result

    result['name'] = info["name"]
    result['count'] = info["count"]
    result['size'] = info["size"]
    result['screenshots'] = info["screenshots"]

    # 收集截图 URL，行内并行下载（6 线程）
    screenshot_urls = [
        s.get("screenshot") for s in info.get("screenshots", [])
        if s.get("screenshot")
    ]
    result['temp_files'] = download_screenshots_parallel(screenshot_urls)

    return result


# ============================================================
# 进度管理
# ============================================================
class ProgressManager:
    """管理处理进度，支持断点续传"""

    def __init__(self, input_dir="Input", output_dir="Output", state_file=".progress.json"):
        self.input_dir = Path(input_dir)
        self.output_dir = Path(output_dir)
        self.state_file = self.output_dir / state_file
        self.state = {}
        self.load_state()

    def load_state(self):
        if self.state_file.exists():
            try:
                with open(self.state_file, 'r', encoding='utf-8') as f:
                    self.state = json.load(f)
                print(f"加载进度状态: 已处理 {self.get_total_processed()} 行")
            except Exception as e:
                print(f"加载状态文件失败: {e}")
                self.state = {}
        else:
            self.state = {}

    def save_state(self):
        self.state['_updated_at'] = datetime.now().isoformat()
        try:
            with open(self.state_file, 'w', encoding='utf-8') as f:
                json.dump(self.state, f, indent=2, ensure_ascii=False)
        except Exception as e:
            print(f"保存状态文件失败: {e}")

    def get_file_state(self, file_path):
        file_key = str(file_path)
        if file_key not in self.state:
            self.state[file_key] = {
                'processed_rows': 0,
                'total_rows': 0,
                'file_hash': '',
                'completed': False,
                'last_modified': '',
                'batches_created': []
            }
        return self.state[file_key]

    def update_file_state(self, file_path, processed_rows, total_rows=None,
                          completed=False, batch_num=None):
        state = self.get_file_state(file_path)
        state['processed_rows'] = processed_rows
        if total_rows is not None:
            state['total_rows'] = total_rows
        if completed:
            state['completed'] = True
        state['last_modified'] = datetime.now().isoformat()
        if batch_num is not None:
            if 'batches_created' not in state:
                state['batches_created'] = []
            if batch_num not in state['batches_created']:
                state['batches_created'].append(batch_num)
        self.save_state()

    def mark_file_processed(self, file_path):
        state = self.get_file_state(file_path)
        state['completed'] = True
        state['last_modified'] = datetime.now().isoformat()
        self.save_state()

    def is_file_completed(self, file_path):
        state = self.get_file_state(file_path)
        return state.get('completed', False)

    def get_processed_rows(self, file_path):
        state = self.get_file_state(file_path)
        return state.get('processed_rows', 0)

    def get_total_processed(self):
        total = 0
        for key, state in self.state.items():
            if key.startswith('_'):
                continue
            total += state.get('processed_rows', 0)
        return total

    def reset_file_state(self, file_path):
        file_key = str(file_path)
        if file_key in self.state:
            del self.state[file_key]
            self.save_state()


# ============================================================
# Excel 辅助函数
# ============================================================
def get_excel_files(input_dir):
    if not input_dir.exists():
        print(f"Input 目录不存在: {input_dir}")
        return []
    excel_files = []
    extensions = ['*.xlsx', '*.xlsm', '*.xls']
    for ext in extensions:
        excel_files.extend(glob.glob(str(input_dir / ext)))
    return [Path(f) for f in excel_files]


def _natural_sort_key(path: Path):
    """
    按文件名做自然数字排序：
      File_000.xlsx < File_001.xlsx < File_009.xlsx < File_010.xlsx < File_100.xlsx
    也兼容 1_xxx / 2_xxx / 10_xxx 等命名。

    返回类型统一的元组列表，避免 str 与 int 直接比较时的 TypeError：
      (0, "前缀字符串")  ← 字符串片段
      (1, 数字)          ← 数字片段
    """
    name = path.stem
    parts = re.split(r'(\d+)', name)
    return [
        (1, int(p)) if p.isdigit() else (0, p.lower())
        for p in parts
    ]


def get_unprocessed_files(input_dir, progress_manager):
    """
    返回所有未完成处理的文件，并**按文件名自然数字序**排序，
    保证每次运行都从最小号文件（File_000.xlsx）开始处理。
    """
    all_files = get_excel_files(input_dir)
    unprocessed = []
    for file_path in all_files:
        if progress_manager.is_file_completed(file_path):
            print(f"  跳过已完成文件: {file_path.name}")
            continue

        file_hash = get_file_hash(file_path)
        state = progress_manager.get_file_state(file_path)
        if state.get('file_hash') and state.get('file_hash') != file_hash:
            print(f"  文件 {file_path.name} 已变更，重置进度")
            progress_manager.reset_file_state(file_path)

        state['file_hash'] = file_hash
        unprocessed.append(file_path)

    # ✅ 关键修复：按文件名自然数字序排序（File_000 → File_001 → File_002 → …）
    unprocessed.sort(key=_natural_sort_key)

    # 🔍 打印排序后的顺序，方便一眼确认
    if unprocessed:
        print("  📑 待处理文件顺序:")
        for i, f in enumerate(unprocessed, 1):
            print(f"     {i}. {f.name}")

    return unprocessed


def count_data_rows_in_file(file_path):
    """统计 Excel 文件中的数据行数（不含标题行）"""
    try:
        wb = load_workbook(file_path, read_only=True)
        ws = wb.active
        row_count = sum(1 for _ in ws.iter_rows(min_row=2))
        wb.close()
        return row_count
    except Exception as e:
        print(f"统计行数失败 {file_path.name}: {e}")
        return 0


def get_headers_from_file(file_path):
    """从 Excel 文件读取标题行"""
    try:
        wb = load_workbook(file_path, read_only=True)
        ws = wb.active
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        wb.close()

        if first_row:
            headers = [str(h) if h else "" for h in first_row[:4]]
            if not any(headers):
                return ["链接", "Resource Name", "Number of Files", "Total File Size"]
            return headers
        else:
            return ["链接", "Resource Name", "Number of Files", "Total File Size"]
    except Exception as e:
        print(f"读取标题行失败 {file_path.name}: {e}")
        return ["链接", "Resource Name", "Number of Files", "Total File Size"]


def save_batch_to_file(batch_data, headers, batch_num, output_dir, base_name):
    """
    将一批（最多 BATCH_SIZE=50 行）数据保存为一个独立 Excel 文件。
    文件名格式: {batch_num:03d}_{base_name}.xlsx
    标题行在第 1 行，数据从第 2 行开始。
    """
    output_filename = f"{batch_num:03d}_{base_name}.xlsx"
    output_path = output_dir / output_filename

    print(f"  📦 保存批次 {batch_num} 到: {output_filename}")
    print(f"  📊 数据行数: {len(batch_data)}")

    workbook = xlsxwriter.Workbook(str(output_path))
    worksheet = workbook.add_worksheet()

    worksheet.set_column(0, 3, 20)

    # 标题行（第 1 行，索引 0）
    for col_idx, header in enumerate(headers):
        worksheet.write(0, col_idx, header)

    # 数据行
    for row_idx, row_data in enumerate(batch_data):
        excel_row = row_idx + 1

        worksheet.write(excel_row, 0, row_data['magnet'])
        worksheet.write(excel_row, 1, row_data['name'])
        worksheet.write(excel_row, 2, row_data['count'])
        worksheet.write(excel_row, 3, row_data['size'])

        worksheet.set_row(excel_row, DEFAULT_ROW_HEIGHT)

        if row_data['temp_files']:
            col = 4
            for tmp_path in row_data['temp_files']:
                if os.path.exists(tmp_path):
                    worksheet.set_column(col, col, DEFAULT_COL_WIDTH)
                    worksheet.insert_image(excel_row, col, tmp_path,
                                           {'x_scale': 1, 'y_scale': 1})
                    col += 1

    workbook.close()
    print(f"  ✅ 批次 {batch_num} 保存完成")


# ============================================================
# 单文件处理（跨行多进程 + 行内多线程 + 阶段性保存）
# ============================================================
def process_single_file(file_path, output_dir, progress_manager,
                        file_index=None, total_files=None):
    """
    处理单个 Excel 文件：
      - 每处理 BATCH_SIZE（50）行，立即保存为一个独立 Excel 文件到 Output
      - 立即更新进度并提交到 Git（阶段性保存，不等到全部完成）
      - 跨行用 ProcessPoolExecutor 并行，行内截图用 ThreadPoolExecutor 并行
    返回: (processed_rows, should_continue)
    """
    file_label = f"[{file_index}/{total_files}] " if file_index and total_files else ""
    print(f"{file_label}处理文件: {file_path.name}")

    start_row = progress_manager.get_processed_rows(file_path)
    if start_row > 0:
        print(f"  从第 {start_row + 2} 行继续（已处理 {start_row} 行数据）")

    try:
        wb_in = load_workbook(file_path, read_only=True)
        ws_in = wb_in.active

        headers = get_headers_from_file(file_path)
        print(f"  标题行: {headers}")

        total_rows = count_data_rows_in_file(file_path)
        print(f"  总数据行数: {total_rows}")

        if start_row >= total_rows:
            progress_manager.mark_file_processed(file_path)
            print(f"  ✓ 文件 {file_path.name} 已全部处理完成！")
            wb_in.close()
            return 0, False

        # 收集本次要处理的行（磁力链接 + Excel 行号）
        rows_to_process = []
        data_row_idx = 0
        for row in ws_in.iter_rows(min_row=2, values_only=True):
            if data_row_idx < start_row:
                data_row_idx += 1
                continue
            if not row or not row[0]:
                print(f"  第 {data_row_idx + 2} 行为空，停止收集")
                break
            excel_row_num = data_row_idx + 2
            rows_to_process.append((str(row[0]).strip(), excel_row_num))
            data_row_idx += 1

        wb_in.close()

        if not rows_to_process:
            progress_manager.mark_file_processed(file_path)
            print(f"  ✓ 文件 {file_path.name} 无新数据")
            return 0, False

        print(f"  本次计划处理: {len(rows_to_process)} 行数据")
        print(f"  每 {BATCH_SIZE} 行保存一个 Excel 文件")
        print(f"  跨行并行进程数: {MAX_WORKERS}，行内截图并行线程数: {SCREENSHOT_THREADS}")
        print(f"  API 限速: {API_RATE_LIMIT} 次/分钟")

        start_time = time.time()
        processed_count = 0
        batch_num = (start_row // BATCH_SIZE) + 1
        file_base = file_path.stem
        hit_limit = False

        # 使用进程池并行处理（跨行）
        with ProcessPoolExecutor(max_workers=MAX_WORKERS) as executor:
            for batch_start in range(0, len(rows_to_process), BATCH_SIZE):
                # 时间检查
                elapsed = time.time() - start_time
                if elapsed > MAX_RUN_TIME * 60:
                    print(f"  已达到运行时间限制: {MAX_RUN_TIME} 分钟")
                    hit_limit = True
                    break

                batch = rows_to_process[batch_start:batch_start + BATCH_SIZE]
                print(f"\n  ▶ 批次 {batch_num}（{len(batch)} 行，"
                      f"进度 {batch_start + 1}-{batch_start + len(batch)}/{len(rows_to_process)}）")

                # 提交所有行到进程池
                futures = [executor.submit(_process_row_worker, arg) for arg in batch]

                # 收集结果
                results = []
                for future in as_completed(futures):
                    try:
                        results.append(future.result())
                    except Exception as e:
                        print(f"  ⚠️ 行处理失败: {e}")

                # 按行号排序以保持顺序
                results.sort(key=lambda r: r['row_num'])

                # ===== 立即保存为独立 Excel 文件（阶段性保存）=====
                save_batch_to_file(results, headers, batch_num, output_dir, file_base)

                processed_count += len(results)

                # 立即更新进度
                progress_manager.update_file_state(
                    file_path,
                    start_row + processed_count,
                    total_rows,
                    completed=False,
                    batch_num=batch_num
                )
                print(f"  💾 进度已保存: {start_row + processed_count}/{total_rows}")

                # 立即提交到 Git（阶段性保存）
                git_commit_files(output_dir, batch_num)

                # 清理临时文件
                for r in results:
                    for tmp in r.get('temp_files', []):
                        try:
                            if os.path.exists(tmp):
                                os.unlink(tmp)
                        except Exception:
                            pass

                batch_num += 1

                elapsed = time.time() - start_time
                print(f"  📈 累计处理 {start_row + processed_count}/{total_rows} 行，"
                      f"耗时 {elapsed / 60:.1f} 分钟")

        total_processed = start_row + processed_count
        if total_processed >= total_rows:
            progress_manager.mark_file_processed(file_path)
            print(f"  ✓ 文件 {file_path.name} 处理完成！")
        else:
            print(f"  ⏳ 文件 {file_path.name} 部分完成: {total_processed}/{total_rows}")

        elapsed = time.time() - start_time
        should_continue = (
            processed_count < len(rows_to_process) and
            elapsed < MAX_RUN_TIME * 60 and
            total_processed < total_rows
        )

        if hit_limit:
            should_continue = True

        return processed_count, should_continue

    except Exception as e:
        print(f"  处理文件 {file_path.name} 时出错: {e}")
        import traceback
        traceback.print_exc()
        return 0, False


# ============================================================
# 主函数
# ============================================================
def main():
    print("=" * 60)
    print("磁力链接批量处理器 - GitHub Actions 版本")
    print("  跨行多进程 + 行内多线程 + 全局限速 + 阶段性保存")
    print("=" * 60)
    print(f"运行时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print()

    print("当前配置:")
    print(f"  BATCH_SIZE: {BATCH_SIZE} 行/批次（每批保存一个 Excel 文件）")
    print(f"  MAX_WORKERS: {MAX_WORKERS} 跨行并行进程")
    print(f"  SCREENSHOT_THREADS: {SCREENSHOT_THREADS} 行内截图并行线程")
    print(f"  API_RATE_LIMIT: {API_RATE_LIMIT} 次/分钟（官方限制）")
    print(f"  MAX_RUN_TIME: {MAX_RUN_TIME} 分钟")
    print(f"  MAX_RETRIES: {MAX_RETRIES} 次")
    print(f"  TIMEOUT: {TIMEOUT} 秒")
    print("-" * 60)

    input_dir = Path("Input")
    output_dir = Path("Output")

    input_dir.mkdir(exist_ok=True)
    output_dir.mkdir(exist_ok=True)

    progress_manager = ProgressManager(input_dir, output_dir)

    unprocessed_files = get_unprocessed_files(input_dir, progress_manager)

    if not unprocessed_files:
        print("所有文件已处理完成！")
        return

    print(f"\n找到 {len(unprocessed_files)} 个未处理的文件（按文件名自然序）:")
    for f in unprocessed_files:
        processed = progress_manager.get_processed_rows(f)
        total = count_data_rows_in_file(f)
        print(f"  - {f.name} (已处理 {processed}/{total} 行数据)")
    print()

    total_processed = 0
    files_processed = 0

    for file_path in unprocessed_files:
        processed, should_continue = process_single_file(
            file_path, output_dir, progress_manager,
            file_index=files_processed + 1, total_files=len(unprocessed_files)
        )
        total_processed += processed
        files_processed += 1

        if should_continue:
            print("\n" + "=" * 60)
            print("⚠️ 已达到运行限制，将在下次运行继续")
            print(f"  本次处理: {processed} 行")
            print(f"  累计处理: {progress_manager.get_total_processed()} 行")
            print("=" * 60)
            return

    print("\n" + "=" * 60)
    print("🎉 所有文件处理完成！")
    print(f"  处理文件数: {files_processed}")
    print(f"  总处理行数: {total_processed}")
    print(f"  输出目录: {output_dir}")
    print("=" * 60)


if __name__ == "__main__":
    main()
